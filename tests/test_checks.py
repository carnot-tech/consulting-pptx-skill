"""機械チェックの自己テスト（標準ライブラリのみ）。

指摘を機械チェックに足したら、ここに「直していない版で FAIL が出る」「直した版で出ない」の
両方を1件ずつ足す。直していない版で発火しないチェックは、測れていないのと同じ。

  python3 -m unittest discover -s tests
"""
import re
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
CHECK = ROOT / "scripts" / "check_deck.py"
GOOD = (ROOT / "tests" / "fixtures" / "good_deck.html").read_text(encoding="utf8")


def run(html, *extra):
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "deck.html"
        p.write_text(html, encoding="utf8")
        extra = [str(Path(d) / a[1:]) if a.startswith("@") else a for a in extra]
        for a in extra:
            if a.endswith(".txt"):
                Path(a).write_text("# 架空の禁止語\nサンプル商事\nre:PJ-\\d{3}\n", encoding="utf8")
        r = subprocess.run([sys.executable, str(CHECK), str(p), *extra], capture_output=True, text=True)
    return r.returncode, [l for l in r.stdout.splitlines() if l.startswith("FAIL")]


def warns(html):
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "deck.html"
        p.write_text(html, encoding="utf8")
        r = subprocess.run([sys.executable, str(CHECK), str(p)], capture_output=True, text=True)
    return [l for l in r.stdout.splitlines() if l.startswith("WARN")]


def inject(old, new):
    assert old in GOOD, old
    return GOOD.replace(old, new, 1)


class GoodDeck(unittest.TestCase):
    def test_good_deck_passes(self):
        code, fails = run(GOOD)
        self.assertEqual((code, fails), (0, []))

    def test_templates_pass_in_template_mode(self):
        for f in ("freeform_parts_16x9.html", "freeform_parts_more_16x9.html"):
            r = subprocess.run([sys.executable, str(CHECK), str(ROOT / "templates" / f), "--template"],
                               capture_output=True, text=True)
            self.assertEqual(r.returncode, 0, f + r.stdout[-400:])


class Placeholders(unittest.TestCase):
    def test_body_placeholder_fires(self):
        code, fails = run(inject("承認待ちは平均3日", "Text 1"))
        self.assertEqual(code, 1)
        self.assertTrue(any("本文にテンプレのプレースホルダー" in f for f in fails), fails)

    def test_part_label_chip_fires(self):
        code, fails = run(inject('<div class="date">現状</div>', '<div class="date">パーツ06｜前提→帰結の2カラム</div>'))
        self.assertTrue(any("プレースホルダー" in f for f in fails), fails)

    def test_type_name_title_fires(self):
        code, fails = run(inject("<h1>1部門で2か月試行し、効果を確かめてから全社へ広げる</h1>", "<h1>軸のある表</h1>"))
        self.assertTrue(any("型名のまま" in f for f in fails), fails)


class OneSentencePerBlock(unittest.TestCase):
    def test_two_sentences_in_card_fires(self):
        code, fails = run(inject("• 例外だけ承認者に回す", "• 例外だけ承認者に回す。規程外の申請は差し戻す"))
        self.assertTrue(any("2文以上" in f for f in fails), fails)

    def test_period_inside_brackets_is_ignored(self):
        code, fails = run(inject("• 例外だけ承認者に回す", "• 例外（金額超過。科目不明）だけ承認者に回す"))
        self.assertFalse(any("2文以上" in f for f in fails), fails)

    def test_single_sentence_with_trailing_period_passes(self):
        code, fails = run(inject("• 例外だけ承認者に回す", "• 例外だけ承認者に回す。"))
        self.assertFalse(any("2文以上" in f for f in fails), fails)


class ForbiddenTerms(unittest.TestCase):
    def test_term_in_body_fires_without_echoing_it(self):
        code, fails = run(inject("営業部", "サンプル商事の営業部"), "--forbid", "@terms.txt")
        self.assertTrue(any("禁止語" in f for f in fails), fails)
        self.assertFalse(any("サンプル商事" in f for f in fails), fails)

    def test_term_in_html_comment_fires(self):
        code, fails = run(inject("</main>", "<!-- PJ-123 向けの下書き --></main>"), "--forbid", "@terms.txt")
        self.assertTrue(any("コメント/属性" in f for f in fails), fails)

    def test_no_list_no_check(self):
        code, fails = run(inject("営業部", "サンプル商事の営業部"))
        self.assertEqual(code, 0)


def run_out(html):
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "deck.html"
        p.write_text(html, encoding="utf8")
        r = subprocess.run([sys.executable, str(CHECK), str(p)], capture_output=True, text=True)
    return r.stdout


class Charset(unittest.TestCase):
    def test_missing_charset_fires(self):
        html = re.sub(r"<meta[^>]*charset[^>]*>", "", GOOD, count=1)
        self.assertNotEqual(html, GOOD)
        code, fails = run(html)
        self.assertTrue(any("charset" in f for f in fails), fails)

    def test_good_deck_has_charset(self):
        code, fails = run(GOOD)
        self.assertFalse(any("charset" in f for f in fails), fails)


class TitleConnector(unittest.TestCase):
    T = "<h1>1部門で2か月試行し、効果を確かめてから全社へ広げる</h1>"

    def test_connector_start_warns(self):
        out = run_out(inject(self.T, "<h1>まずは1部門で2か月試行し、効果を確かめてから全社へ広げる</h1>"))
        self.assertIn("接続詞で始まる", out)

    def test_word_starting_with_mata_does_not_warn(self):
        out = run_out(inject(self.T, "<h1>またがる2部門で試行し、効果を確かめてから全社へ広げる</h1>"))
        self.assertNotIn("接続詞で始まる", out)


class ProductionMeta(unittest.TestCase):
    def test_tool_name_fails(self):
        code, fails = run(inject("</main>", "<p>本資料は consulting-pptx-skill で作成</p></main>"))
        self.assertTrue(any("制作メタ" in f for f in fails), fails)

    def test_github_url_as_source_only_warns(self):
        code, fails = run(inject("</main>", "<p>出典: github.com/example/repo</p></main>"))
        self.assertFalse(any("制作メタ" in f for f in fails), fails)


def _has_playwright():
    return (ROOT / "node_modules" / "playwright").exists()


@unittest.skipUnless(_has_playwright(), "npm run setup で playwright を入れると実行される")
class EmptyArea(unittest.TestCase):
    def layout(self, html):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "deck.html"
            p.write_text(html, encoding="utf8")
            r = subprocess.run(["node", str(ROOT / "scripts" / "check_layout.mjs"), str(p)],
                               capture_output=True, text=True, cwd=ROOT)
        return r.returncode, r.stdout

    def test_good_deck_has_no_empty_page(self):
        code, out = self.layout(GOOD)
        self.assertEqual(code, 0, out)

    def test_half_empty_page_fires(self):
        html = re.sub(r'\s*<tr><td class="ax">2026年7月</td>.*?変わらず</td></tr>', "", GOOD, flags=re.S)
        html = html.replace("\n          <li>1人あたりの申請件数が2倍になる</li>\n          <li>承認の遅れが立替者の不満になる</li>", "")
        self.assertNotEqual(html, GOOD)
        code, out = self.layout(html)
        self.assertEqual(code, 1, out)
        self.assertIn("p3: 版面の", out)


if __name__ == "__main__":
    unittest.main()


class DashAndHeads(unittest.TestCase):
    """§6 非該当「—」は連結ではない／§4.49 基本パーツ集のカラム見出し .colh も対象"""
    def test_label_dash_join_warns(self):
        w = warns(inject("承認待ちは平均3日", "承認待ち — 平均3日"))
        self.assertTrue(any("ダッシュ" in l for l in w), w)

    def test_standalone_na_dash_is_ignored(self):
        w = warns(inject("承認待ちは平均3日", "承認待ちは平均3日</p><p>—</p><p>—（関与しない）"))
        self.assertFalse(any("ダッシュ" in l for l in w), w)

    def test_colh_conjunction_head_warns(self):
        w = warns(inject("承認待ちは平均3日", '承認待ちは平均3日<div class="colh">だから、全社展開を急ぐ</div>'))
        self.assertTrue(any("接続詞" in l for l in w), w)


def _has_pptx():
    try:
        import pptx  # noqa: F401
        return True
    except ImportError:
        return False


@unittest.skipUnless(_has_pptx(), "python-pptx が入っていると実行される")
class PptxLayout(unittest.TestCase):
    """check_deck_layout.py: 図形に隠れた文字・中身の無い箱（線は箱と取り違えない）"""

    def _deck(self, build):
        from pptx import Presentation
        from pptx.util import Inches
        from pptx.enum.shapes import MSO_SHAPE
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        s = prs.slides.add_slide(prs.slide_layouts[6])
        build(s, Inches, MSO_SHAPE)
        d = tempfile.mkdtemp()
        p = Path(d) / "deck.pptx"
        prs.save(p)
        sys.path.insert(0, str(ROOT / "scripts"))
        import check_deck_layout
        return [k for _, k, _ in check_deck_layout.check(str(p))]

    @staticmethod
    def _solid(shape):
        shape.fill.solid()
        return shape

    def test_text_under_opaque_box_fires(self):
        def build(s, In, M):
            tb = s.shapes.add_textbox(In(1), In(1), In(4), In(1))
            tb.text_frame.text = "隠れている文字"
            self._solid(s.shapes.add_shape(M.RECTANGLE, In(0.9), In(0.9), In(4.2), In(1.2)))
        self.assertIn("hidden", self._deck(build))

    def test_visible_text_and_thin_line_do_not_fire(self):
        def build(s, In, M):
            tb = s.shapes.add_textbox(In(1), In(1), In(4), In(1))
            tb.text_frame.text = "見えている文字"
            self._solid(s.shapes.add_shape(M.RECTANGLE, In(1), In(3), In(1.3), In(0.01)))
        kinds = self._deck(build)
        self.assertNotIn("hidden", kinds)
        self.assertNotIn("empty", kinds)

    def test_empty_filled_box_fires(self):
        def build(s, In, M):
            self._solid(s.shapes.add_shape(M.RECTANGLE, In(1), In(1), In(2), In(1)))
        self.assertIn("empty", self._deck(build))
